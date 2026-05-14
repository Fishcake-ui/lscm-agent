"""
build_customers.py — Convert Combined_Master_Abir_JamesSmith.xlsx to customers.json
Run this once locally, and again any time you update the master file.

Install: pip install openpyxl

Usage:
  python build_customers.py
  python build_customers.py --input path/to/Combined_Master.xlsx --output customers.json
"""

import argparse
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)


NOISE = re.compile(
    r'\b(pvt|private|limited|ltd|llp|inc|co|company|the|hotel|hotels|'
    r'restaurant|restaurants|services|hospitality|enterprises|industries)\b',
    re.IGNORECASE,
)
PUNCT = re.compile(r'[\(\)\[\]\.,&/\-_]+')


def gen_aliases(dest, abir_tally, js_tally):
    """Generate fuzzy-match aliases from destination + Tally names."""
    raw = [s for s in [dest, abir_tally, js_tally] if s]
    out = set()
    for s in raw:
        s = str(s).strip()
        out.add(s.lower())                              # full lowercase
        head = s.split(',')[0].strip()                  # before first comma
        out.add(head.lower())
        cleaned = PUNCT.sub(' ', NOISE.sub(' ', head))
        cleaned = re.sub(r'\s+', ' ', cleaned).strip().lower()
        if cleaned:
            out.add(cleaned)
        words = [w for w in cleaned.split() if len(w) > 1]
        if len(words) >= 2:
            out.add(' '.join(words[:2]))               # first 2 significant words
        if words:
            out.add(words[0])                           # first significant word

    out.discard(dest.lower() if dest else '')           # the name itself isn't an alias
    out = {a for a in out if a and len(a) >= 2 and not a.isdigit()}
    return sorted(out)


def build(input_path: str, output_path: str):
    if not os.path.exists(input_path):
        print(f"ERROR: input file not found: {input_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(input_path, data_only=True)
    if 'Master (Unified)' not in wb.sheetnames:
        print(f"ERROR: expected sheet 'Master (Unified)' not found. Sheets: {wb.sheetnames}")
        sys.exit(1)

    ws = wb['Master (Unified)']
    # row 1 = title, row 2 = column headers, row 3+ = data
    headers = [c.value for c in ws[2]]
    idx = {h: i for i, h in enumerate(headers)}

    required = ['Delivery Destination', 'Area / Locality', 'Pincode', 'Match Status',
                'Abir Tally Name', 'James Smith Tally Name', 'Address', 'Phone']
    missing = [r for r in required if r not in idx]
    if missing:
        print(f"ERROR: master sheet missing columns: {missing}")
        sys.exit(1)

    customers = []
    ab_ct = js_ct = both_ct = 0

    for r in range(3, ws.max_row + 1):
        row = [ws.cell(row=r, column=i + 1).value for i in range(len(headers))]
        dest = row[idx['Delivery Destination']]
        if not dest:
            continue

        status = row[idx['Match Status']] or ''
        area = row[idx['Area / Locality']] or ''
        pincode = row[idx['Pincode']]
        abir_tally = row[idx['Abir Tally Name']]
        js_tally = row[idx['James Smith Tally Name']]
        address = row[idx['Address']] or ''
        phone = row[idx['Phone']] or ''

        if 'Abir Only' in status:
            ab_ct += 1
            cust_id = f'AF-{ab_ct:03d}'
            company = 'Abir Foods'
        elif 'James Smith Only' in status:
            js_ct += 1
            cust_id = f'JS-{js_ct:03d}'
            company = 'James Smith'
        else:
            both_ct += 1
            cust_id = f'BOTH-{both_ct:03d}'
            company = 'Both'

        customers.append({
            'id': cust_id,
            'name': str(dest),
            'aliases': gen_aliases(dest, abir_tally, js_tally),
            'company': company,
            'area': str(area),
            'pincode': str(pincode) if pincode else '',
            'abir_tally': abir_tally,
            'js_tally': js_tally,
            'address': str(address),
            'phone': phone if phone and phone != '—' else None,
            'credit_days': 30,
            'match_status': status,
        })

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump({
            'customers': customers,
            'count': len(customers),
            'source': os.path.basename(input_path),
        }, f, indent=2, ensure_ascii=False)

    print(f"Wrote {len(customers)} customers to {output_path}")
    print(f"  Abir Only:    {ab_ct}")
    print(f"  James Smith:  {js_ct}")
    print(f"  Both:         {both_ct}")


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('--input', default='Combined_Master_Abir_JamesSmith.xlsx',
                   help='Path to the combined master xlsx (default: ./Combined_Master_Abir_JamesSmith.xlsx)')
    p.add_argument('--output', default='customers.json',
                   help='Path to write customers.json (default: ./customers.json)')
    args = p.parse_args()
    build(args.input, args.output)
